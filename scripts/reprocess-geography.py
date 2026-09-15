#!/usr/bin/env python3
"""Reprocess the two approved CIT sources without writing a database.
Usage: python scripts/reprocess-geography.py /secure/source-folder output-folder
Requires scripts/requirements-data.txt. Hash mismatch fails closed.
"""
import csv
import hashlib
import io
import json
import pathlib
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from openpyxl import load_workbook
from replication import fingerprint, COLS, BASE
HASHES={'COD_MUNICIPAL.zip':'a5947915a7213cddde00682a51d0734ea6b6ec2307d237d1e7937edde6766b99','CEP5.xlsx':'74ad34907a4ee418ededa872d3530cc11ae89270ed666331a6879ea72cb8cf13'}
UF=dict(zip(['11','12','13','14','15','16','17','21','22','23','24','25','26','27','28','29','31','32','33','35','41','42','43','50','51','52','53'],['RO','AC','AM','RR','PA','AP','TO','MA','PI','CE','RN','PB','PE','AL','SE','BA','MG','ES','RJ','SP','PR','SC','RS','MS','MT','GO','DF']))
ALIASES={('RR','São Luiz'):'1400605',('RN','Arês'):'2401206',('RN','Açu'):'2400208'}
NS={'t':'urn:oasis:names:tc:opendocument:xmlns:table:1.0','x':'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
def ods_rows(data):
 with zipfile.ZipFile(io.BytesIO(data)) as z:
  if z.getinfo('content.xml').file_size>50_000_000: raise ValueError('Unexpected ODS size')
  root=ET.fromstring(z.read('content.xml'))
 for row in root.findall('.//t:table-row',NS):
  values=[]
  for cell in row.findall('t:table-cell',NS):
   value=' '.join(''.join(p.itertext()) for p in cell.findall('x:p',NS))
   repeated=int(cell.attrib.get('{'+NS['t']+'}number-columns-repeated','1'))
   values.extend([value]*min(repeated,30))
   if len(values)>=30: break
  if len(values)>8 and re.fullmatch('[0-9]{7}',values[7]): yield values

def main():
 source=pathlib.Path(sys.argv[1]); target=pathlib.Path(sys.argv[2]);target.mkdir(parents=True,exist_ok=True)
 for file,expected in HASHES.items():
  if hashlib.sha256((source/file).read_bytes()).hexdigest()!=expected:raise ValueError('Unapproved source hash: '+file)
 result={t:[] for t in COLS}
 with zipfile.ZipFile(source/'COD_MUNICIPAL.zip') as z:
  for suffix,table in [('MUNICIPIOS','dim_municipio'),('DISTRITOS','dim_distrito'),('SUBDISTRITOS','dim_subdistrito')]:
   names=[n for n in z.namelist() if n.endswith('_'+suffix+'.ods') and 'RELATORIO' in n]
   if len(names)!=1: raise ValueError('Expected exactly one territorial report')
   for v in ods_rows(z.read(names[0])):
    common={'ano_dtb':2025,'data_base_dtb':'2025-12-31','ativo':True}
    if table=='dim_municipio':data=dict(zip(['cod_uf','nome_uf','cod_regiao_intermediaria','regiao_intermediaria','cod_regiao_imediata','regiao_imediata','cod_municipio_dtb','cod_municipal','municipio'],v[:9]))
    elif table=='dim_distrito':data={'cod_distrito':v[10],'cod_municipal':v[7],'distrito_dtb':v[9],'distrito':v[11]}
    else:data={'cod_subdistrito':v[13],'cod_distrito':v[10],'cod_municipal':v[7],'subdistrito_dtb':v[12],'subdistrito':v[14]}
    result[table].append(data|common)
 lookup={(UF[m['cod_uf']],m['municipio']):m['cod_municipal'] for m in result['dim_municipio']}
 workbook=load_workbook(source/'CEP5.xlsx',read_only=True,data_only=True)
 sheet=workbook['Resultados'];headers=next(sheet.iter_rows(values_only=True))
 positions={str(x):i for i,x in enumerate(headers)}
 for row in sheet.iter_rows(min_row=2,values_only=True):
  cp=str(row[positions['CEP5']]);municipio=row[positions['Município']];uf=row[positions['Estado']]
  if not re.fullmatch('[0-9]{5}',cp):raise ValueError('CEP5 must already have five digits')
  code=lookup.get((uf,municipio)) or ALIASES.get((uf,municipio))
  if not code:raise ValueError('Unresolved municipality in approved source')
  result['dim_cep5'].append({'cod_municipal':code,'cep5':cp,'municipio_origem':municipio,'uf_origem':uf,'metodo_resolucao':'ALIAS_HOMOLOGADO' if (uf,municipio) in ALIASES else 'EXATO','ativo':True})
 workbook.close()
 report={}
 for table,rows in result.items():
  fp=fingerprint(table,rows);expected=BASE[table]['source_primary_replica_sha256']
  if fp!=expected or len(rows)!=BASE[table]['source_rows']:raise ValueError('Canonical content mismatch: '+table)
  with (target/(table+'.csv')).open('w',encoding='utf-8',newline='') as f:
   w=csv.DictWriter(f,fieldnames=COLS[table]);w.writeheader();w.writerows({k:str(v).lower() if isinstance(v,bool) else v for k,v in row.items()} for row in rows)
  report[table]={'rows':len(rows),'sha256':fp}
 (target/'verification.json').write_text(json.dumps(report,indent=2)+'\n')
 print(json.dumps(report,indent=2))
if __name__=='__main__':main()
